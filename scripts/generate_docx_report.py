import sys
import json
import os
import docx

def generate_docx(input_json_path, output_docx_path):
    with open(input_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    template_path = r'D:\ChiLinh\BaoCaoDanhGia\template.docx'
    if not os.path.exists(template_path):
        template_path = os.path.join(os.path.dirname(__file__), '..', 'template.docx')

    doc = docx.Document(template_path)

    student = data.get('student', {})
    session = data.get('session', {})
    stats = data.get('stats', {})
    repertoires = data.get('repertoires', [])

    # Fill Table 0 (Admin info)
    t0 = doc.tables[0]
    t0.rows[0].cells[0].text = f"Full name / Họ và tên: {student.get('full_name', '')} ({student.get('nickname', '')})"
    gender_vn = 'Nam' if student.get('gender') == 'male' else ('Nữ' if student.get('gender') == 'female' else 'Khác')
    t0.rows[0].cells[1].text = f"Sex / Giới tính: {gender_vn}"
    
    t0.rows[1].cells[0].text = f"Date of birth / Ngày sinh: {student.get('dob', '')}"
    age_str = data.get('age_str', '')
    t0.rows[1].cells[1].text = f"Age / Tuổi: {age_str}"
    
    current_class = student.get('current_class', 'Inclusive')
    t0.rows[2].cells[0].text = f"Current grade / Lớp hiện tại: {current_class}"
    t0.rows[2].cells[1].text = f"Current grade / Lớp hiện tại: {current_class}"

    test_date = session.get('assessment_date', student.get('evaluation_date', ''))
    t0.rows[4].cells[0].text = f"Dates of testing / Ngày thực hiện: {test_date} (Đợt {session.get('round_number', 1)})"
    t0.rows[4].cells[1].text = f"Dates of testing / Ngày thực hiện: {test_date} (Đợt {session.get('round_number', 1)})"

    assessor = session.get('assessor_name', student.get('evaluator_name', 'Ngô Thị Thùy Linh'))
    t0.rows[5].cells[0].text = f"Examinators / Người thực hiện: {assessor}"
    t0.rows[5].cells[1].text = f"Examinators / Người thực hiện: {assessor}"

    # Fill Table 1 (25 repertoires)
    if len(doc.tables) > 1:
        t1 = doc.tables[1]
        rep_map = {r['code']: r for r in repertoires}

        for row in t1.rows[1:]:
            code_cell = row.cells[0].text.strip()
            if code_cell in rep_map:
                rep_data = rep_map[code_cell]
                # Completion ratio
                row.cells[2].text = f"{rep_data.get('percentage', 0)}% ({rep_data.get('earned', 0)}/{rep_data.get('total', 0)}đ)"
                # Norm benchmark
                row.cells[3].text = f"{rep_data.get('normBenchmark', 0)}% (Chuẩn tuổi)"
                # Evaluation comment
                pct = rep_data.get('percentage', 0)
                if pct >= 80:
                    eval_text = "Kỹ năng thành thục tốt, đạt mức phát triển mong đợi."
                elif pct >= 40:
                    eval_text = "Đang phát triển, cần tiếp tục hỗ trợ và duy trì luyện tập."
                else:
                    eval_text = "Kỹ năng còn hạn chế, cần ưu tiên đưa vào mục tiêu can thiệp."
                row.cells[4].text = eval_text

    # Update summary text paragraphs
    brief_text = f"Tổng số điểm đạt được: {stats.get('totalEarned', 0)}/{stats.get('totalMax', 1478)} điểm ({stats.get('totalPercent', 0)}%). Đã lượng giá {stats.get('scoredSkillsCount', 0)}/544 bài kỹ năng trong Đợt {session.get('round_number', 1)}."
    
    general_summary = session.get('general_summary', '') or "Trẻ có biểu hiện hợp tác trong quá trình đánh giá. Cần tập trung hỗ trợ các lĩnh vực giao tiếp chức năng và tương tác xã hội."
    strengths = session.get('strengths', '') or "Trẻ có khả năng tập trung tốt với các hoạt động yêu thích và phản ứng tích cực với phần thưởng."
    
    iep_goals = session.get('iep_goals', [])
    goals_text = ""
    if iep_goals:
        for idx, g in enumerate(iep_goals, 1):
            if isinstance(g, dict):
                goals_text += f"{idx}. [{g.get('repertoireCode', '')}] {g.get('title', '')}: {g.get('targetCriteria', '')}\n"
            else:
                goals_text += f"{idx}. {g}\n"
    else:
        goals_text = "1. Phát triển khả năng hiểu và phản hồi ngôn ngữ tiếp nhận.\n2. Gia tăng tần suất tương tác và giao tiếp bằng mắt.\n3. Hỗ trợ các kỹ năng tự lập cơ bản."

    for p in doc.paragraphs:
        if 'BRIEF RESULT' in p.text:
            p.text = f"BRIEF RESULT / KẾT QUẢ NGẮN GỌN\n{brief_text}"
        elif 'GENERAL ASSESSMENT OF CURRENT COMPETENCIES' in p.text:
            p.text = f"GENERAL ASSESSMENT OF CURRENT COMPETENCIES / ĐÁNH GIÁ TỔNG QUAN NĂNG LỰC HIỆN TẠI\n{general_summary}"
        elif 'HIGHLIGHTS IN THE EVALUATION PROCESS' in p.text:
            p.text = f"HIGHLIGHTS IN THE EVALUATION PROCESS / ĐIỂM NỔI BẬT TRONG QUÁ TRÌNH ĐÁNH GIÁ\n{strengths}"
        elif 'PRIORITY INTERVENTION GOALS' in p.text:
            p.text = f"PRIORITY INTERVENTION GOALS (for the next 6 – 12 months) / MỤC TIÊU ƯU TIÊN CAN THIỆP (trong 6 – 12 tháng tới)\n{goals_text}"

    doc.save(output_docx_path)
    print(f"Docx generated successfully: {output_docx_path}")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_docx_report.py <input_json> <output_docx>")
        sys.exit(1)
    generate_docx(sys.argv[1], sys.argv[2])
