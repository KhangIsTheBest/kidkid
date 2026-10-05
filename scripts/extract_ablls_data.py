import openpyxl
import json
import os
import re
import sys

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return " ".join(text.split())

def extract_ablls():
    excel_path = r"D:\ChiLinh\DanhGiaTongHop\ABLLS_DanhGia.xlsm"
    print(f"Loading master workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 1. Map Age files to detect age milestones for each skill
    age_skills = {2: set(), 3: set(), 4: set(), 5: set()}
    for age in [2, 3, 4, 5]:
        age_file = rf"D:\ChiLinh\DanhGiaTheoTuoi\ABLLS_{age}Tuoi.xlsm"
        if os.path.exists(age_file):
            wb_age = openpyxl.load_workbook(age_file, data_only=True)
            for sname in wb_age.sheetnames:
                ws = wb_age[sname]
                for row in ws.iter_rows(values_only=True):
                    if row and row[0]:
                        code = str(row[0]).strip()
                        if re.match(r'^[A-Z]\d+$', code):
                            age_skills[age].add(code)
            print(f"Age {age} loaded: {len(age_skills[age])} skills")

    # 2. Extract Norms (Mốc tuổi chuẩn)
    norms_data = {}
    norms_intervals = []
    if "Mốc tuổi chuẩn (Norms)" in wb.sheetnames:
        ws_norms = wb["Mốc tuổi chuẩn (Norms)"]
        rows = list(ws_norms.iter_rows(values_only=True))
        
        # Row 2 contains column headers
        header_row = rows[2]
        # Intervals are columns 2 to 20 (up to 60T)
        for col_idx in range(2, len(header_row)):
            val = header_row[col_idx]
            if val is not None and "T (" in str(val):
                norms_intervals.append(str(val).strip())

        for row in rows[3:]:
            if not row or not row[0]:
                continue
            rep_name = clean_text(row[0])
            rep_code = clean_text(row[1]) if len(row) > 1 and row[1] else ""
            if not rep_name or rep_name.startswith("="):
                continue
            
            values = []
            for col_idx in range(2, 2 + len(norms_intervals)):
                if col_idx < len(row):
                    v = row[col_idx]
                    try:
                        values.append(round(float(v), 4) if v is not None else 0.0)
                    except (ValueError, TypeError):
                        values.append(0.0)
                else:
                    values.append(0.0)
            
            key = rep_code if rep_code else rep_name
            norms_data[key] = {
                "name": rep_name,
                "code": rep_code,
                "values": values
            }
        print(f"Norms extracted for {len(norms_data)} categories across {len(norms_intervals)} intervals: {norms_intervals[0]} -> {norms_intervals[-1]}")

    # 3. Extract Repertoires & Skills
    repertoire_meta = [
        {"code": "A", "sheet": "HỢP TÁC & KHEN THƯỞNG", "name": "Hợp tác & Khen thưởng", "name_en": "Cooperation & Reinforcer Effectiveness"},
        {"code": "B", "sheet": "KHẢ NĂNG QUAN SÁT", "name": "Khả năng quan sát", "name_en": "Visual Performance"},
        {"code": "C", "sheet": "NHẬN BIẾT", "name": "Nhận biết (Hiểu ngôn ngữ)", "name_en": "Receptive Language"},
        {"code": "D", "sheet": "BẮT CHƯỚC", "name": "Bắt chước", "name_en": "Motor Imitation"},
        {"code": "E", "sheet": "NÓI THEO", "name": "Nói theo", "name_en": "Vocal Imitation (Echoic)"},
        {"code": "F", "sheet": "YÊU CẦU", "name": "Yêu cầu", "name_en": "Requests (Mands)"},
        {"code": "G", "sheet": "GỌI TÊN", "name": "Gọi tên", "name_en": "Labeling (Tacts)"},
        {"code": "H", "sheet": "HỘI THOẠI LIÊN TƯỞNG", "name": "Hội thoại liên tưởng", "name_en": "Intraverbals"},
        {"code": "I", "sheet": "NGÔN NGỮ TỰ PHÁT", "name": "Ngôn ngữ tự phát", "name_en": "Spontaneous Vocalizations"},
        {"code": "J", "sheet": "CÚ PHÁP & NGỮ PHÁP", "name": "Cú pháp & Ngữ pháp", "name_en": "Syntax & Grammar"},
        {"code": "K", "sheet": "VUI CHƠI", "name": "Vui chơi & Giải trí", "name_en": "Play & Leisure"},
        {"code": "L", "sheet": "TƯƠNG TÁC XÃ HỘI", "name": "Tương tác xã hội", "name_en": "Social Interaction"},
        {"code": "M", "sheet": "HỌC NHÓM", "name": "Học nhóm", "name_en": "Group Instruction"},
        {"code": "N", "sheet": "TUÂN THEO NỀ NẾP LỚP", "name": "Tuân theo nề nếp lớp", "name_en": "Classroom Routines"},
        {"code": "P", "sheet": "TỔNG HỢP KHÁI QUÁT", "name": "Tổng hợp khái quát", "name_en": "Generalized Responding"},
        {"code": "Q", "sheet": "ĐỌC", "name": "Đọc", "name_en": "Reading"},
        {"code": "R", "sheet": "TOÁN", "name": "Toán", "name_en": "Math"},
        {"code": "S", "sheet": "VIẾT", "name": "Viết", "name_en": "Writing"},
        {"code": "T", "sheet": "GHÉP VẦN", "name": "Ghép vần", "name_en": "Spelling"},
        {"code": "U", "sheet": "TỰ MẶC QUẦN ÁO", "name": "Tự mặc quần áo", "name_en": "Dressing"},
        {"code": "V", "sheet": "ĂN", "name": "Ăn uống", "name_en": "Eating"},
        {"code": "W", "sheet": "TỰ CHĂM SÓC", "name": "Tự chăm sóc vệ sinh", "name_en": "Grooming"},
        {"code": "X", "sheet": "ĐI VỆ SINH", "name": "Đi vệ sinh", "name_en": "Toileting"},
        {"code": "Y", "sheet": "VẬN ĐỘNG THÔ", "name": "Vận động thô", "name_en": "Gross Motor"},
        {"code": "Z", "sheet": "VẬN ĐỘNG TINH", "name": "Vận động tinh", "name_en": "Fine Motor"}
    ]

    repertoires_output = []
    all_skills_output = []

    for rep in repertoire_meta:
        sname = rep["sheet"]
        if sname not in wb.sheetnames:
            print(f"Warning: Sheet {sname} not found!")
            continue
        ws = wb[sname]
        
        rep_skills = []
        current_skill = None

        for row in ws.iter_rows(values_only=True):
            if not row:
                continue
            first_cell = clean_text(row[0]) if len(row) > 0 else ""
            
            # Check if this row starts a new skill (e.g. A1, A2, B1...)
            if re.match(r'^[A-Z]\d+$', first_cell):
                if current_skill:
                    rep_skills.append(current_skill)
                    all_skills_output.append(current_skill)
                
                code = first_cell
                title = clean_text(row[1]) if len(row) > 1 else ""
                objective = clean_text(row[2]) if len(row) > 2 else ""
                question = clean_text(row[3]) if len(row) > 3 else ""
                example = clean_text(row[4]) if len(row) > 4 else ""
                
                score_val = row[5] if len(row) > 5 else 0
                try:
                    score_num = int(score_val) if score_val is not None else 0
                except (ValueError, TypeError):
                    score_num = 0
                
                criteria_text = clean_text(row[6]) if len(row) > 6 else ""
                
                # Determine ages this skill belongs to
                ages_included = [age for age in [2, 3, 4, 5] if code in age_skills[age]]

                current_skill = {
                    "code": code,
                    "repertoire_code": rep["code"],
                    "repertoire_name": rep["name"],
                    "title": title,
                    "objective": objective,
                    "question": question,
                    "example": example,
                    "max_score": score_num,
                    "criteria": [],
                    "ages": ages_included
                }
                if score_num > 0 and criteria_text:
                    current_skill["criteria"].append({
                        "score": score_num,
                        "description": criteria_text
                    })
            elif current_skill:
                # Sub-row with criteria for lower scores (e.g., 3, 2, 1)
                score_val = row[5] if len(row) > 5 else None
                criteria_text = clean_text(row[6]) if len(row) > 6 else ""
                if score_val is not None and str(score_val).strip():
                    try:
                        score_num = int(score_val)
                        if score_num > current_skill["max_score"]:
                            current_skill["max_score"] = score_num
                        current_skill["criteria"].append({
                            "score": score_num,
                            "description": criteria_text
                        })
                    except (ValueError, TypeError):
                        pass

        if current_skill:
            rep_skills.append(current_skill)
            all_skills_output.append(current_skill)

        # Sort criteria by score ascending
        for s in rep_skills:
            s["criteria"].sort(key=lambda c: c["score"])

        rep_item = {
            "code": rep["code"],
            "name": rep["name"],
            "name_en": rep["name_en"],
            "skills_count": len(rep_skills),
            "total_points": sum(s["max_score"] for s in rep_skills),
            "skills": [s["code"] for s in rep_skills]
        }
        repertoires_output.append(rep_item)
        print(f"Repertoire {rep['code']} - {rep['name']}: {len(rep_skills)} skills, total {rep_item['total_points']} pts")

    final_payload = {
        "framework": "ABLLS-R",
        "name": "Assessment of Basic Language and Learning Skills - Revised",
        "version": "Partington 2010",
        "total_skills": len(all_skills_output),
        "total_repertoires": len(repertoires_output),
        "norms_intervals": norms_intervals,
        "norms": norms_data,
        "repertoires": repertoires_output,
        "skills": all_skills_output
    }

    out_dir = r"E:\Projects\kidkid\data"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "ablls_master_dataset.json")
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(final_payload, f, ensure_ascii=False, indent=2)

    print(f"\nSUCCESS: Extracted {len(all_skills_output)} skills into {out_file}")

if __name__ == "__main__":
    extract_ablls()
