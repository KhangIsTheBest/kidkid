import sys
import json
import os
import openpyxl
from openpyxl.cell.cell import Cell

# Grid column ranges for 25 repertoires in sheet 'Tổng hợp' and 'Mốc X tuổi'
REP_GRID_CONFIG = {
    'A': {'start_col': 2, 'max_items': 19},
    'B': {'start_col': 8, 'max_items': 27},
    'C': {'start_col': 14, 'max_items': 57},
    'D': {'start_col': 20, 'max_items': 27},
    'E': {'start_col': 26, 'max_items': 20},
    'F': {'start_col': 32, 'max_items': 29},
    'G': {'start_col': 38, 'max_items': 47},
    'H': {'start_col': 44, 'max_items': 49},
    'I': {'start_col': 50, 'max_items': 9},
    'J': {'start_col': 56, 'max_items': 20},
    'K': {'start_col': 62, 'max_items': 15},
    'L': {'start_col': 68, 'max_items': 34},
    'M': {'start_col': 74, 'max_items': 12},
    'N': {'start_col': 80, 'max_items': 10},
    'P': {'start_col': 86, 'max_items': 6},
    'Q': {'start_col': 92, 'max_items': 17},
    'R': {'start_col': 98, 'max_items': 29},
    'S': {'start_col': 104, 'max_items': 10},
    'T': {'start_col': 110, 'max_items': 7},
    'U': {'start_col': 116, 'max_items': 15},
    'V': {'start_col': 122, 'max_items': 10},
    'W': {'start_col': 128, 'max_items': 7},
    'X': {'start_col': 134, 'max_items': 10},
    'Y': {'start_col': 140, 'max_items': 30},
    'Z': {'start_col': 146, 'max_items': 28},
}

def generate_xlsm(input_json_path, output_xlsm_path):
    with open(input_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    template_path = r'D:\ChiLinh\DanhGiaTongHop\ABLLS_DanhGia.xlsm'
    if not os.path.exists(template_path):
        template_path = os.path.join(os.path.dirname(__file__), '..', 'ABLLS_DanhGia.xlsm')

    wb = openpyxl.load_workbook(template_path, keep_vba=True)

    student = data.get('student', {})
    sessions = data.get('sessions', [])
    # Find latest session or map by round number
    session_map = {s['round_number']: s for s in sessions}
    latest_session = sessions[-1] if sessions else {}

    scores = latest_session.get('scores', {})
    skill_rounds = latest_session.get('skill_rounds', {})

    # Sheets to update student header info
    summary_sheets = ['Tổng hợp', 'Mốc 2 tuổi', 'Mốc 3 tuổi', 'Mốc 4 tuổi', 'Mốc 5 tuổi']

    for s_name in summary_sheets:
        if s_name not in wb.sheetnames:
            continue
        ws = wb[s_name]

        # Student Name and DOB (top-left cells of merged ranges B4:J5 and B6:J7)
        nickname_str = f" ({student.get('nickname')})" if student.get('nickname') else ""
        ws['B4'] = f"Họ tên : {student.get('full_name', '')}{nickname_str}"
        ws['B6'] = f"Ngày sinh: {student.get('dob', '')}"

        # Evaluator and testing dates per round
        for r_num in [1, 2, 3, 4]:
            s_sess = session_map.get(r_num)
            r_idx = 9 + r_num  # Row 10 for R1, 11 for R2, 12 for R3, 13 for R4
            if s_sess:
                ws.cell(r_idx, 2).value = s_sess.get('assessor_name', 'Giáo viên can thiệp')
                ws.cell(r_idx, 5).value = s_sess.get('assessment_date', '')

        # Fill Tower Grid cells
        # Read skills metadata from data['skills']
        skills = data.get('skills', [])
        # Group by repertoire
        rep_skills = {}
        for s in skills:
            rep_skills.setdefault(s['repertoire_code'], []).append(s)

        for rep_code, config in REP_GRID_CONFIG.items():
            s_list = rep_skills.get(rep_code, [])
            s_list.sort(key=lambda x: int(''.join(filter(str.isdigit, x['code'])) or '0'))

            for item_idx, skill in enumerate(s_list, 1):
                row = 60 - (item_idx - 1)
                start_col = config['start_col']
                max_score = skill.get('max_score', 1)
                cur_score = scores.get(skill['code'], 0)
                skill_round_map = skill_rounds.get(skill['code'], {})

                for lvl in range(1, max_score + 1):
                    if max_score == 4:
                        offset = lvl - 1
                    elif max_score == 2:
                        offset = (lvl - 1) * 2
                    elif max_score == 1:
                        offset = 0
                    else:
                        offset = lvl - 1

                    col = start_col + offset
                    target_cell = ws.cell(row, col)
                    # Only assign value if target_cell is an active Cell, not a read-only MergedCell
                    if isinstance(target_cell, Cell):
                        if cur_score >= lvl:
                            rnd = skill_round_map.get(str(lvl), skill_round_map.get(lvl, latest_session.get('round_number', 1)))
                            target_cell.value = int(rnd)
                        else:
                            target_cell.value = None

    wb.save(output_xlsm_path)
    print(f"Xlsm generated successfully: {output_xlsm_path}")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_xlsm_report.py <input_json> <output_xlsm>")
        sys.exit(1)
    generate_xlsm(sys.argv[1], sys.argv[2])
